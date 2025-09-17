

from flask import Flask, request, jsonify
from flask_cors import CORS
import os
from typing import Dict, List, Any
import json
import PyPDF2
import docx
from io import BytesIO
import logging
from tenacity import retry, stop_after_attempt, wait_exponential, retry_if_exception_type
from dotenv import load_dotenv
import hashlib
import ollama

# Load environment variables
load_dotenv()

app = Flask(__name__)
CORS(app)

# Configure logging
logging.basicConfig(level=logging.DEBUG)

class CVAnalyzer:
    def __init__(self):
        self.model = "gemma2:9b-instruct-q4_K_M"
        self.use_mock = os.getenv('USE_MOCK', 'False').lower() == 'true'

    def extract_text_from_file(self, file_content: bytes, filename: str) -> str:
        """Extract text from uploaded file based on file type"""
        try:
            if filename.lower().endswith('.pdf'):
                return self._extract_from_pdf(file_content)
            elif filename.lower().endswith(('.doc', '.docx')):
                return self._extract_from_docx(file_content)
            elif filename.lower().endswith('.txt'):
                return file_content.decode('utf-8')
            else:
                raise ValueError(f"Unsupported file type: {filename}")
        except Exception as e:
            raise Exception(f"Error extracting text from {filename}: {str(e)}")

    def _extract_from_pdf(self, file_content: bytes) -> str:
        """Extract text from PDF file"""
        pdf_reader = PyPDF2.PdfReader(BytesIO(file_content))
        text = ""
        for page in pdf_reader.pages:
            extracted = page.extract_text()
            if extracted:
                text += extracted + "\n"
        return text

    def _extract_from_docx(self, file_content: bytes) -> str:
        """Extract text from DOCX file"""
        doc = docx.Document(BytesIO(file_content))
        text = ""
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        return text

    def _mock_analysis(self) -> Dict[str, Any]:
        """Mock response for testing without API"""
        return {
            "overall_score": 85,
            "requirements_score": 90,
            "wishes_score": 80,
            "requirements": [
                {
                    "id": "req1",
                    "type": "require",
                    "title": "SQL Proficiency",
                    "description": "Advanced SQL skills",
                    "match": True,
                    "percentage": 95,
                    "explanation": "CV shows 5+ years of SQL experience in data analysis projects."
                }
            ],
            "wishes": [
                {
                    "id": "wish1",
                    "type": "wish",
                    "title": "Power BI Experience",
                    "description": "Familiarity with Power BI",
                    "match": True,
                    "percentage": 80,
                    "explanation": "CV mentions dashboard creation with Power BI."
                }
            ]
        }

    @retry(
        stop=stop_after_attempt(5),
        wait=wait_exponential(multiplier=2, min=4, max=60),
        retry=retry_if_exception_type((json.JSONDecodeError, ollama.ResponseError))
    )
    def _call_ollama(self, prompt: str, system_content: str, temperature: float = 0.3) -> str:
        """Internal API call to Ollama with retry"""
        try:
            response = ollama.chat(
                model=self.model,
                messages=[
                    {"role": "system", "content": system_content},
                    {"role": "user", "content": prompt}
                ],
                options={"temperature": temperature}
            )
            content = response['message']['content'].strip()
            logging.debug(f"Raw Ollama response: {content}")
            return content
        except ollama.ResponseError as e:
            logging.error(f"Ollama API error: {str(e)}")
            raise
        except Exception as e:
            logging.error(f"Unexpected error: {str(e)}")
            raise

    def analyze_cv_assignment_match(self, cv_text: str, assignment_text: str) -> Dict[str, Any]:
        """Analyze CV against assignment requirements using gemma2:9b-instruct-q4_K_M"""
        if self.use_mock:
            logging.debug("Using mock analysis response")
            return self._mock_analysis()

        # Validate assignment_text
        if not assignment_text or assignment_text.isspace():
            logging.error("Assignment requirements are empty")
            return {
                "overall_score": 0,
                "requirements_score": 0,
                "wishes_score": 0,
                "requirements": [],
                "wishes": [],
                "error": "Assignment requirements cannot be empty"
            }

        # Truncate inputs to reduce tokens
        cv_text = cv_text[:1000]
        assignment_text = assignment_text[:1000]

        # Cache key based on inputs
        cache_key = hashlib.md5((cv_text + assignment_text).encode()).hexdigest()
        cache_file = f"cache/{cache_key}.json"
        if os.path.exists(cache_file):
            logging.debug(f"Returning cached result for key: {cache_key}")
            with open(cache_file, 'r') as f:
                return json.load(f)

        prompt = f"""
        You are an expert HR consultant specializing in data professional recruitment.
        Analyze the provided CV against the assignment requirements and return a response in **valid JSON format only**.
        Do not include any text outside the JSON object, such as explanations, <think> tags, or code fences.

        ASSIGNMENT REQUIREMENTS:
        {assignment_text}

        CONSULTANT CV:
        {cv_text}

        Return a JSON object with the following structure:
        {{
            "overall_score": <number 0-100>,
            "requirements_score": <number 0-100>,
            "wishes_score": <number 0-100>,
            "requirements": [
                {{
                    "id": "<unique_id>",
                    "type": "require",
                    "title": "<requirement_title>",
                    "description": "<requirement_description>",
                    "match": <boolean>,
                    "percentage": <number 0-100>,
                    "explanation": "<detailed_explanation>"
                }}
            ],
            "wishes": [
                {{
                    "id": "<unique_id>",
                    "type": "wish",
                    "title": "<wish_title>",
                    "description": "<wish_description>",
                    "match": <boolean>,
                    "percentage": <number 0-100>,
                    "explanation": "<detailed_explanation>"
                }}
            ]
        }}

        Focus on:
        - Technical skills matching (e.g., Azure, SQL, Power BI)
        - Experience relevance
        - Education requirements
        - Soft skills and cultural fit
        - Provide specific examples from the CV

        If no requirements are provided, return an empty requirements and wishes list with scores of 0.
        Ensure the response is a valid JSON object with no additional text, <think> tags, or code fences.
        """

        try:
            content = self._call_ollama(
                prompt,
                "You are an expert HR consultant. Return only valid JSON with no additional text, <think> tags, or code fences.",
                temperature=0.3
            )
            result = json.loads(content)

            # Cache result
            os.makedirs("cache", exist_ok=True)
            with open(cache_file, 'w') as f:
                json.dump(result, f)
            logging.debug(f"Cached result for key: {cache_key}")

            return result
        except json.JSONDecodeError as e:
            logging.error(f"JSON parsing error: {e}, raw content: {content}")
            return {
                "overall_score": 0,
                "requirements_score": 0,
                "wishes_score": 0,
                "requirements": [],
                "wishes": [],
                "error": f"Invalid JSON response from Ollama: {str(e)}"
            }
        except ollama.ResponseError as e:
            logging.error(f"Ollama error: {str(e)}")
            return {
                "overall_score": 0,
                "requirements_score": 0,
                "wishes_score": 0,
                "requirements": [],
                "wishes": [],
                "error": f"Ollama API error: {str(e)}"
            }
        except Exception as e:
            logging.error(f"API error: {str(e)}")
            return {
                "overall_score": 0,
                "requirements_score": 0,
                "wishes_score": 0,
                "requirements": [],
                "wishes": [],
                "error": f"Error in CV analysis: {str(e)}"
            }

    def generate_motivations(self, cv_text: str, requirements: List[Dict], consultant_name: str) -> Dict[str, str]:
        """Generate personalized motivations for each requirement"""
        if self.use_mock:
            return {req['id']: "Mock motivation for testing." for req in requirements}

        motivations = {}
        cv_text = cv_text[:1000]  # Truncate

        for req in requirements:
            prompt = f"""
            You are an expert HR consultant tasked with writing highly professional and concise motivations for a job candidate.
            Do not include any <think> tags, internal processing notes, or code fences in the response.

            CONSULTANT: {consultant_name}
            REQUIREMENT: {req['title']} - {req['description']}
            MATCH PERCENTAGE: {req['percentage']}%
            CV CONTENT:
            {cv_text}

            Generate a motivation statement that:
            - Is 2-3 sentences long (50-75 words)
            - Uses a formal, professional tone
            - Directly addresses the requirement with specific examples from the CV
            - Highlights relevant skills, experience, or achievements
            - If the match percentage is below 70%, acknowledges gaps but emphasizes transferable skills or eagerness to learn
            - Avoids generic phrases and focuses on concrete, tailored content
            - Returns plain text with no additional formatting, <think> tags, or code fences

            Example:
            My five years of experience developing complex SQL queries for data warehousing projects directly aligns with the requirement for advanced SQL proficiency. I have successfully optimized database performance for large-scale analytics, as demonstrated in my recent project at [Company Name]. This expertise ensures I can effectively contribute to your data management initiatives.
            """

            try:
                content = self._call_ollama(
                    prompt,
                    "You are an expert HR consultant. Return only plain text with no additional formatting, <think> tags, or code fences.",
                    temperature=0.3
                )
                motivations[req['id']] = content.strip()
            except ollama.ResponseError as e:
                motivations[req['id']] = f"Ollama API error: {str(e)}"
            except Exception as e:
                motivations[req['id']] = f"Error generating motivation: {str(e)}"

        return motivations

    def generate_cover_letter(self, cv_text: str, assignment_info: Dict, consultant_name: str, analysis_result: Dict) -> str:
        """Generate a personalized cover letter"""
        if self.use_mock:
            return "Mock cover letter for testing."

        cv_text = cv_text[:1000]  # Reduced truncation
        prompt = f"""
        You are an expert HR consultant tasked with writing a highly professional cover letter for a data professional.
        Do not include any <think> tags, internal processing notes, or code fences in the response.

        CONSULTANT: {consultant_name}
        CLIENT: {assignment_info.get('client', 'Hiring Manager')}
        POSITION: {assignment_info.get('title', 'Data Professional')}
        OVERALL MATCH SCORE: {analysis_result.get('overall_score', 0)}%
        CV CONTENT:
        {cv_text}
        ASSIGNMENT DESCRIPTION:
        {assignment_info.get('description', '')[:1000]}

        Generate a cover letter that:
        - Is 300-400 words, structured in 4-5 paragraphs
        - Uses a formal, professional, and enthusiastic tone
        - Starts with a personalized introduction addressing the client or hiring manager
        - Highlights 2-3 key qualifications from the CV that align with the assignment
        - Addresses potential gaps (e.g., missing certifications) with transferable skills or eagerness to learn
        - Demonstrates knowledge of the client’s industry or needs
        - Ends with a strong call to action and professional closing
        - Uses proper formatting (e.g., Dear [Client], paragraphs, Sincerely, [Name])
        - Returns plain text with no additional formatting, <think> tags, or code fences

        Example structure:
        Dear [Client],

        I am excited to apply for the [Position] role at [Client]. [1-2 sentences about interest and fit.]

        [Paragraph on key qualification 1 with specific CV example.]
        [Paragraph on key qualification 2 with specific CV example.]
        [Optional paragraph addressing gaps or additional relevant skills.]

        I am eager to contribute to [Client]’s success and discuss how my skills align with your needs. Thank you for considering my application. I look forward to the opportunity to discuss further.

        Sincerely,
        {consultant_name}
        """

        try:
            content = self._call_ollama(
                prompt,
                "You are an expert HR consultant. Return only plain text with no additional formatting, <think> tags, or code fences.",
                temperature=0.4
            )
            return content.strip()
        except ollama.ResponseError as e:
            return f"Ollama API error: {str(e)}"
        except Exception as e:
            return f"Error generating cover letter: {str(e)}"

    def generate_introduction_email(self, consultant_info: Dict, assignment_info: Dict, analysis_result: Dict) -> str:
        """Generate a professional introduction email"""
        if self.use_mock:
            return "Mock email for testing."

        consultant_name = consultant_info.get('name', 'the consultant')
        client_name = assignment_info.get('client', 'the client')
        position_title = assignment_info.get('title', 'the position')
        contact_person = consultant_info.get('contactPerson', 'Hiring Manager')
        overall_score = analysis_result.get('overall_score', 0)

        prompt = f"""
        You are an expert staffing agency representative tasked with writing a professional introduction email.
        Do not include any <think> tags, internal processing notes, or code fences in the response.

        CONSULTANT: {consultant_name}
        CLIENT: {client_name}
        POSITION: {position_title}
        CONTACT PERSON: {contact_person}
        MATCH SCORE: {overall_score}%

        Generate an email that:
        - Has a subject line: "Introducing {consultant_name} for {position_title} at {client_name}"
        - Addresses the contact person formally (e.g., Dear [Contact Person])
        - Introduces the consultant with 2-3 key strengths from their CV
        - Mentions the high match score and alignment with the role
        - Includes a standard terms of offer section with placeholders (e.g., [Rate], [Start Date])
        - Lists attachments (e.g., CV, Motivation, Cover Letter)
        - Ends with a professional closing and contact information for Wanita Bajnath
        - Is 200-300 words, concise, and professional
        - Returns plain text with no additional formatting, <think> tags, or code fences

        Example structure:
        Subject: Introducing {consultant_name} for {position_title} at {client_name}

        Dear {contact_person},

        I am pleased to introduce {consultant_name}, an exceptional candidate for the {position_title} role at {client_name}. [1-2 sentences on key strengths and match score.]

        [Brief paragraph on consultant’s relevant experience and skills.]

        Terms of Offer:
        - Rate: [Rate]
        - Availability: [Start Date]
        - Contract Duration: [Duration]

        Attached are {consultant_name}’s CV, motivation, and cover letter for your review. Please contact me to discuss next steps or arrange an interview.

        Best regards,
        Wanita Bajnath
        Senior Staffing Consultant
        [Agency Name]
        Email: wanita.bajnath@[agency].com
        Phone: [Phone Number]
        """

        try:
            content = self._call_ollama(
                prompt,
                "You are an expert staffing agency representative. Return only plain text with no additional formatting, <think> tags, or code fences.",
                temperature=0.3
            )
            return content.strip()
        except ollama.ResponseError as e:
            return f"Ollama API error: {str(e)}"
        except Exception as e:
            return f"Error generating email: {str(e)}"

# Initialize analyzer
analyzer = CVAnalyzer()

@app.route('/api/analyze', methods=['POST'])
def analyze_cv():
    """Analyze CV against assignment requirements"""
    try:
        cv_file = request.files.get('cv_file')
        assignment_file = request.files.get('assignment_file')
        assignment_data = json.loads(request.form.get('assignment_data', '{}'))
        consultant_data = json.loads(request.form.get('consultant_data', '{}'))
        
        if not cv_file:
            return jsonify({'error': 'CV file is required'}), 400
        
        cv_content = cv_file.read()
        cv_text = analyzer.extract_text_from_file(cv_content, cv_file.filename)
        
        assignment_text = assignment_data.get('description', '')
        if assignment_file:
            assignment_content = assignment_file.read()
            assignment_file_text = analyzer.extract_text_from_file(assignment_content, assignment_file.filename)
            assignment_text = f"{assignment_text}\n\n{assignment_file_text}"
        
        analysis_result = analyzer.analyze_cv_assignment_match(cv_text, assignment_text)
        
        return jsonify({
            'success': True,
            'analysis': analysis_result
        })
    
    except Exception as e:
        logging.error(f"API /api/analyze error: {e}")
        return jsonify({'error': str(e)}), 500

@app.route('/api/generate-motivations', methods=['POST'])
def generate_motivations():
    """Generate motivations for requirements"""
    try:
        data = request.json
        cv_text = data.get('cv_text')
        requirements = data.get('requirements', [])
        consultant_name = data.get('consultant_name')
        
        motivations = analyzer.generate_motivations(cv_text, requirements, consultant_name)
        
        return jsonify({
            'success': True,
            'motivations': motivations
        })
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/generate-cover-letter', methods=['POST'])
def generate_cover_letter():
    """Generate cover letter"""
    try:
        data = request.json
        cv_text = data.get('cv_text')
        assignment_info = data.get('assignment_info')
        consultant_name = data.get('consultant_name')
        analysis_result = data.get('analysis_result')
        
        cover_letter = analyzer.generate_cover_letter(cv_text, assignment_info, consultant_name, analysis_result)
        
        return jsonify({
            'success': True,
            'cover_letter': cover_letter
        })
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/generate-email', methods=['POST'])
def generate_email():
    """Generate introduction email"""
    try:
        data = request.json
        consultant_info = data.get('consultant_info')
        assignment_info = data.get('assignment_info')
        analysis_result = data.get('analysis_result')
        
        email = analyzer.generate_introduction_email(consultant_info, assignment_info, analysis_result)
        
        return jsonify({
            'success': True,
            'email': email
        })
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/health', methods=['GET'])
def health_check():
    """Health check endpoint"""
    return jsonify({'status': 'healthy', 'service': 'CV Analysis API'})

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)




